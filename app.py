
import os
import json
import re
import io
import calendar
from datetime import datetime, date
from pathlib import Path
from flask import Flask, jsonify, send_from_directory, send_file, request, Response

PRICE_SPREADSHEET_ID = "1l1qub-I2zuLKLDP2RJFGiDNTIBuGEAxI7PTxIDmfYi4"
SCOPES = ["https://www.googleapis.com/auth/spreadsheets.readonly"]

DATA_FILE = "search_data.json"
EXCEPTION_FILE = "exceptions.json"
AUTO_SYNC_MINUTES = 10
VALID_CATEGORIES = ["식품", "뷰티", "생활", "의류", "패션", "잡화", "신발", "도서", "아동", "가전", "미분류"]

app = Flask(__name__, static_folder=".")


def clean_text(v):
    if v is None:
        return ""
    t = str(v).strip()
    if t.lower() in ["nan", "none", "nat"]:
        return ""
    return t


def to_number(v):
    t = clean_text(v).replace(",", "").replace("원", "")
    try:
        return int(float(t))
    except Exception:
        return 0



def normalize_box_no(value):
    """
    박스번호를 검색 가능한 형태로 정리한다.
    예: '57번', 'BOX 57', '57.0' -> '57'
    숫자가 아닌 박스명은 원문을 유지한다.
    """
    text = clean_text(value)
    if not text:
        return ""

    numeric = re.fullmatch(r"\s*(\d+)(?:\.0+)?\s*", text)
    if numeric:
        return str(int(numeric.group(1)))

    labeled = re.search(r"(?:box|박스)?\s*0*(\d+)\s*번?", text, flags=re.I)
    if labeled:
        return str(int(labeled.group(1)))

    return text

def sheet_year(title):
    m = re.search(r"(20\d{2})", str(title))
    return int(m.group(1)) if m else date.today().year


def normalize_date_text(text, default_year=None):
    text = clean_text(text)
    if not text:
        return ""

    m = re.search(r"(?<!\d)(20\d{2})(\d{2})(\d{2})(?!\d)", text)
    if m:
        y, mo, d = map(int, m.groups())
        try:
            return date(y, mo, d).strftime("%Y-%m-%d")
        except Exception:
            return ""

    m = re.search(r"(20\d{2})[-./년\s]+(\d{1,2})[-./월\s]+(\d{1,2})", text)
    if m:
        y, mo, d = map(int, m.groups())
        try:
            return date(y, mo, d).strftime("%Y-%m-%d")
        except Exception:
            return ""

    m = re.search(r"(?<!\d)(\d{2})[-./](\d{1,2})[-./](\d{1,2})(?!\d)", text)
    if m:
        y, mo, d = map(int, m.groups())
        try:
            return date(y + 2000, mo, d).strftime("%Y-%m-%d")
        except Exception:
            return ""

    m = re.search(r"(\d{1,2})\s*월\s*(\d{1,2})\s*일", text)
    if m and default_year:
        mo, d = map(int, m.groups())
        try:
            return date(int(default_year), mo, d).strftime("%Y-%m-%d")
        except Exception:
            return ""

    return ""


def detect_inbound_date_row(row, default_year):
    for c in row:
        t = clean_text(c)
        if not t:
            continue
        if re.search(r"\d{1,2}\s*월\s*\d{1,2}\s*일", t):
            return normalize_date_text(t, default_year=default_year)
        if re.search(r"20\d{2}[-./]\d{1,2}[-./]\d{1,2}", t):
            return normalize_date_text(t, default_year=default_year)
    return ""


def normalize_header(t):
    return clean_text(t).replace("\n", "").replace(" ", "")


def is_header_row(row):
    joined = "|".join([normalize_header(c) for c in row])
    has_product = ("상품명" in joined) or ("품명" in joined)
    has_structure = any(k in joined for k in ["박스번호", "박스", "형태", "대분류", "가격", "금액", "판매가"])
    return has_product and has_structure


def find_col(headers, names):
    h = [normalize_header(x) for x in headers]
    for name in names:
        n = name.replace(" ", "").replace("\n", "")
        for i, v in enumerate(h):
            if n in v:
                return i
    return None


def cell(row, idx):
    if idx is None or idx >= len(row):
        return ""
    return clean_text(row[idx])


def extract_category(text):
    text = str(text or "")
    for c in VALID_CATEGORIES:
        if c != "미분류" and c in text:
            return c
    return "미분류"


def extract_price_band(text):
    text = clean_text(text)
    if not text:
        return ""
    m = re.search(r"(?<!\d)(\d{1,3})[\.,](0)(?!\d)", text)
    if m:
        return f"{int(m.group(1))}.0"
    amounts = re.findall(r"(?<!\d)(\d{1,3}(?:,\d{3})+|\d{4,6})(?!\d)", text)
    for a in amounts:
        n = to_number(a)
        if 1000 <= n <= 300000:
            return f"{round(n / 1000, 1):.1f}"
    return ""



def extract_price_amounts(text):
    """
    가격 셀 안의 금액 후보를 추출한다.
    사이즈 150, 160 같은 숫자는 제외하고 1,000원 이상 금액만 사용한다.
    """
    text = clean_text(text)
    if not text:
        return []

    found = re.findall(r"(?<!\d)(\d{1,3}(?:,\d{3})+|\d{4,6})(?!\d)", text)
    amounts = []
    for raw in found:
        try:
            value = int(raw.replace(",", ""))
        except Exception:
            continue
        if 1000 <= value <= 999999:
            amounts.append(value)

    # 원문 순서를 유지하며 중복 제거
    unique = []
    for value in amounts:
        if value not in unique:
            unique.append(value)
    return unique


def build_price_info(price_raw, note=""):
    """
    가격을 숫자 하나로 억지 변환하지 않고 원문과 표시 방식을 함께 반환한다.

    반환:
    - price: 단일가격일 때만 숫자, 그 외 0
    - price_type: single / range / discretion / multiple / attached / text / missing
    - price_display: 카드에 표시할 짧은 문구
    - price_raw: 시트의 가격 원문
    - price_amounts: 원문에서 추출한 금액 목록
    - show_detail: 가격 상세 펼침 표시 여부
    """
    raw = clean_text(price_raw)
    note_text = clean_text(note)
    combined = f"{raw}\n{note_text}".lower()
    amounts = extract_price_amounts(raw)

    if not raw:
        return {
            "price": 0,
            "price_type": "missing",
            "price_display": "가격 정보 없음",
            "price_raw": "",
            "price_amounts": [],
            "show_detail": False,
        }

    # 숫자 하나만 적힌 단일 가격
    numeric_text = raw.replace(",", "").replace("원", "").strip()
    if re.fullmatch(r"\d+(?:\.0+)?", numeric_text):
        try:
            value = int(float(numeric_text))
        except Exception:
            value = 0
        return {
            "price": value,
            "price_type": "single",
            "price_display": f"{value:,}원" if value else raw,
            "price_raw": raw,
            "price_amounts": [value] if value else [],
            "show_detail": False,
        }

    # 상품에 부착된 가격표를 확인해야 하는 경우
    if "부착된 가격" in raw or "부착 가격" in raw:
        return {
            "price": 0,
            "price_type": "attached",
            "price_display": "상품 부착가 확인",
            "price_raw": raw,
            "price_amounts": amounts,
            "show_detail": False,
        }

    # 매장/권역 재량 가격
    if "재량" in combined:
        if len(amounts) >= 2:
            price_display = f"매장·권역 재량 · {min(amounts):,}~{max(amounts):,}원"
        elif len(amounts) == 1:
            price_display = f"매장·권역 재량 · {amounts[0]:,}원"
        else:
            price_display = "매장·권역 재량 가격"
        return {
            "price": 0,
            "price_type": "discretion",
            "price_display": price_display,
            "price_raw": raw,
            "price_amounts": amounts,
            "show_detail": True,
        }

    # 줄바꿈과 상품별 설명이 있는 복합 가격
    category_words = [
        "슬리퍼", "샌들", "운동화", "구두", "워커", "로퍼", "가방",
        "의류", "잡화", "아동", "성인", "티셔츠", "바지", "셔츠",
        "블라우스", "원피스", "자켓", "점퍼", "사이즈", "size"
    ]
    has_category_lines = "\n" in raw and any(word in raw.lower() for word in category_words)

    if has_category_lines or ("\n" in raw and len(amounts) >= 2):
        return {
            "price": 0,
            "price_type": "multiple",
            "price_display": "종류별 가격 · 상세보기",
            "price_raw": raw,
            "price_amounts": amounts,
            "show_detail": True,
        }

    # 1,000 ~ 3,000 형태의 범위 가격
    if len(amounts) >= 2 and any(mark in raw for mark in ["~", "～", "-", "–", "—"]):
        return {
            "price": 0,
            "price_type": "range",
            "price_display": f"{min(amounts):,}~{max(amounts):,}원",
            "price_raw": raw,
            "price_amounts": amounts,
            "show_detail": True,
        }

    # 3,000 / 5,000 / 7,000 형태
    if len(amounts) >= 2:
        if len(amounts) <= 4:
            display = " / ".join(f"{value:,}" for value in amounts) + "원"
        else:
            display = f"복수가격 · {min(amounts):,}~{max(amounts):,}원"
        return {
            "price": 0,
            "price_type": "multiple",
            "price_display": display,
            "price_raw": raw,
            "price_amounts": amounts,
            "show_detail": True,
        }

    # 숫자 추출은 안 되지만 유효한 가격 안내 문구가 있는 경우
    first_line = raw.splitlines()[0].strip()
    short_text = first_line if len(first_line) <= 30 else first_line[:30] + "…"
    return {
        "price": 0,
        "price_type": "text",
        "price_display": short_text or "가격 상세 확인",
        "price_raw": raw,
        "price_amounts": amounts,
        "show_detail": len(raw.splitlines()) > 1 or len(raw) > 30,
    }

def section_by_expiry(expiry):
    if not expiry:
        return "", "기한확인필요"
    try:
        exp = datetime.strptime(expiry, "%Y-%m-%d").date()
    except Exception:
        return "", "기한확인필요"

    days = (exp - date.today()).days
    if days < 0:
        sec = "만료"
    elif days <= 7:
        sec = "7일이내"
    elif days <= 30:
        sec = "30일이내"
    elif days <= 60:
        sec = "60일이내"
    elif days <= 90:
        sec = "90일이내"
    elif days >= 180:
        sec = "장기재고"
    else:
        sec = "전체"
    return days, sec



def load_exceptions():
    path = Path(EXCEPTION_FILE)
    if not path.exists():
        return {}
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def save_exceptions(data):
    with open(EXCEPTION_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def get_exception_mode(product_name, exceptions):
    name = clean_text(product_name)
    if not name:
        return ""
    for key in sorted(exceptions.keys(), key=len, reverse=True):
        key_clean = clean_text(key)
        if key_clean and key_clean in name:
            return exceptions.get(key, {}).get("mode", "")
    return ""



def make_item_key(store, inbound, box_no, product, sheet, row):
    """
    같은 상품이라도 매장/입고일/박스/행이 다르면 다른 물건으로 봄.
    매장명이 없으면 시트+행까지 포함해 충돌을 줄임.
    """
    parts = [
        clean_text(store) or "전체",
        clean_text(inbound),
        clean_text(box_no),
        clean_text(product),
        clean_text(sheet),
        str(row or "")
    ]
    return "|".join(parts)


def find_item_override(item_key, product_name, exceptions):
    """
    1순위: item_key 정확히 일치
    2순위: 기존 상품명 예외 부분일치
    """
    if item_key and item_key in exceptions:
        return item_key, exceptions[item_key]

    name = clean_text(product_name)
    for key in sorted(exceptions.keys(), key=len, reverse=True):
        key_clean = clean_text(key)
        if key_clean and key_clean in name:
            return key, exceptions.get(key, {})

    return "", {}


def parse_sheet(values, title, exceptions=None):
    y = sheet_year(title)
    exceptions = exceptions or {}
    rows = []
    current_inbound = ""
    current_box_no = ""
    headers = None
    col = {}

    for r_idx, row in enumerate(values, start=1):
        if not any(clean_text(c) for c in row):
            continue

        if is_header_row(row):
            current_box_no = ""
            headers = row
            col = {
                "box": find_col(row, ["박스번호", "박스"]),
                "type": find_col(row, ["형태", "구분"]),
                "category": find_col(row, ["대분류", "분류", "카테고리"]),
                "product": find_col(row, ["상품명", "품명"]),
                "barcode": find_col(row, ["바코드"]),
                "price": find_col(row, ["가격", "금액", "판매가", "단가"]),
                "note": find_col(row, ["비고", "메모", "참고"]),
                "store": find_col(row, ["매장"]),
                "expiry": find_col(row, ["유통기한", "소비기한", "기한"]),
                "inbound": find_col(row, ["입고일", "입고일자"]),
            }
            continue

        date_row = detect_inbound_date_row(row, y)
        if date_row:
            current_inbound = date_row
            current_box_no = ""
            continue

        if not headers or not col:
            continue

        product = cell(row, col.get("product"))
        if not product:
            continue

        if "가격 이외의 사항은" in product or "바코드 사용" in product or "문의 부탁" in product:
            continue

        category_raw = cell(row, col.get("category"))
        category = category_raw if category_raw in VALID_CATEGORIES else extract_category(product)
        item_type = cell(row, col.get("type"))
        raw_box_no = cell(row, col.get("box"))
        normalized_box_no = normalize_box_no(raw_box_no)
        if normalized_box_no:
            current_box_no = normalized_box_no
        box_no = current_box_no
        barcode = cell(row, col.get("barcode"))
        price_raw = cell(row, col.get("price"))
        note = cell(row, col.get("note"))
        store = cell(row, col.get("store"))

        inbound = normalize_date_text(cell(row, col.get("inbound")), y) or current_inbound
        expiry = normalize_date_text(cell(row, col.get("expiry")), y)
        if not expiry:
            expiry = normalize_date_text(f"{product} {note}", y)

        price_info = build_price_info(price_raw, note)
        price = price_info["price"]
        price_type = price_info["price_type"]
        price_display = price_info["price_display"]
        price_amounts = price_info["price_amounts"]
        show_price_detail = price_info["show_detail"]

        band = extract_price_band(price_raw) or extract_price_band(product) or extract_price_band(note)
        if not band and price:
            band = f"{round(price / 1000, 1):.1f}"

        days, sec = section_by_expiry(expiry)

        if category == "식품":
            manage = "유통기한관리" if expiry else "기한확인필요"
        else:
            manage = "유통기한관리" if expiry else "비관리대상"

        exception_mode = get_exception_mode(product, exceptions)
        # 예외/수정 설정 적용
        item_key = make_item_key(store, inbound, box_no, product, title, r_idx)
        matched_key, override = find_item_override(item_key, product, exceptions)
        exception_mode = clean_text(override.get("mode", "")) if override else ""

        override_expiry = clean_text(override.get("expiryDate", "")) if override else ""
        override_note = clean_text(override.get("memo", "")) if override else ""

        if override_expiry:
            expiry = override_expiry
            days, sec = section_by_expiry(expiry)
            manage = "유통기한관리"
            exception_mode = "expiry_override"
        elif exception_mode == "manufacture":
            manage = "기한확인필요"
            sec = "기한확인필요"
            days = ""
        elif exception_mode == "exclude":
            manage = "비관리대상"
            sec = "기한관리제외"
            days = ""
        elif exception_mode == "expiry":
            pass

        search = f"{title} {inbound} {category_raw} {category} {product} {price} {band} {price_raw} {price_display} {expiry} {note} {barcode} {box_no} {store} {item_type} {exception_mode} {override_note}".lower()

        rows.append({
            "manage": manage,
            "section": sec,
            "category": category,
            "majorCategory": category_raw,
            "product": product,
            "price": price,
            "priceBand": band,
            "priceRaw": price_raw,
            "priceType": price_type,
            "priceDisplay": price_display,
            "priceAmounts": price_amounts,
            "showPriceDetail": show_price_detail,
            "inboundDate": inbound,
            "expiryDate": expiry,
            "daysLeft": days,
            "note": note,
            "barcode": barcode,
            "boxNo": box_no,
            "type": item_type,
            "store": store,
            "sheet": title,
            "row": r_idx,
            "search": search,
            "exceptionMode": exception_mode,
            "itemKey": item_key,
            "overrideKey": matched_key,
            "overrideMemo": override_note,
        })

    return rows


def get_credentials():
    import gspread
    from google.oauth2.service_account import Credentials

    json_text = os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON", "").strip()

    if json_text:
        info = json.loads(json_text)
        credentials = Credentials.from_service_account_info(info, scopes=SCOPES)
        return gspread.authorize(credentials)

    # 로컬 테스트용
    local_path = Path(r"C:\GPAM\gpam.json")
    if local_path.exists():
        credentials = Credentials.from_service_account_file(str(local_path), scopes=SCOPES)
        return gspread.authorize(credentials)

    raise FileNotFoundError("GOOGLE_SERVICE_ACCOUNT_JSON 환경변수가 없고, C:\\GPAM\\gpam.json도 없습니다.")


def sync_data():
    client = get_credentials()
    spreadsheet = client.open_by_key(PRICE_SPREADSHEET_ID)

    all_rows = []
    loaded = []
    failed = []
    exceptions = load_exceptions()

    for ws in spreadsheet.worksheets():
        title = ws.title
        if not ("★2026년★" in title or "★2025년★" in title):
            continue

        try:
            parsed = parse_sheet(ws.get_all_values(), title, exceptions)
            if parsed:
                all_rows.extend(parsed)
                loaded.append(title)
            else:
                failed.append(title)
        except Exception as e:
            failed.append(f"{title}: {e}")

    all_rows.sort(key=lambda x: (x.get("inboundDate") or "", x.get("sheet") or "", x.get("row") or 0), reverse=True)

    output = {
        "updatedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "count": len(all_rows),
        "loadedSheets": loaded,
        "failedSheets": failed,
        "items": all_rows,
    }

    with open(DATA_FILE, "w", encoding="utf-8") as f:
        json.dump(output, f, ensure_ascii=False)

    return output


@app.route("/")
def index():
    # 기존 index.html은 유지하고 월별 유통기한 조회 바로가기만 추가한다.
    index_path = Path(app.static_folder or ".") / "index.html"
    if not index_path.exists():
        return send_from_directory(".", "index.html")

    html = index_path.read_text(encoding="utf-8-sig")
    link = (
        '<a href="/expiry_calendar" aria-label="월별 유통기한 조회" '
        'style="position:fixed;right:18px;bottom:18px;z-index:9999;'
        'padding:12px 18px;border-radius:999px;'
        'background:#147d73;color:#fff;text-decoration:none;'
        'font:600 14px system-ui,sans-serif;'
        'box-shadow:0 4px 16px rgba(0,0,0,.2)">'
        '📅 월별 유통기한 조회</a>'
    )
    if re.search(r"</body\s*>", html, flags=re.I):
        html = re.sub(r"</body\s*>", lambda m: link + m.group(0), html,
                      count=1, flags=re.I)
    else:
        html += link
    return Response(html, mimetype="text/html")


@app.route("/search_data.json")
def data():
    should_sync = not Path(DATA_FILE).exists()
    sync_error = ""

    if not should_sync:
        try:
            with open(DATA_FILE, "r", encoding="utf-8") as f:
                current = json.load(f)
            updated_at = current.get("updatedAt", "")
            dt = datetime.strptime(updated_at, "%Y-%m-%d %H:%M:%S")
            age_minutes = (datetime.now() - dt).total_seconds() / 60
            should_sync = age_minutes >= AUTO_SYNC_MINUTES
        except Exception:
            should_sync = True

    if should_sync:
        try:
            sync_data()
        except Exception as e:
            sync_error = str(e)

    if Path(DATA_FILE).exists():
        response = send_from_directory(".", DATA_FILE)
        response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        return response

    return jsonify({
        "ok": False,
        "message": sync_error or "가격리스트 데이터를 생성하지 못했습니다."
    }), 503


@app.route("/sync")
def sync():
    try:
        result = sync_data()
        return jsonify({
            "ok": True,
            "message": "동기화 완료",
            "data": {
                "updatedAt": result["updatedAt"],
                "count": result["count"],
                "loadedSheets": result["loadedSheets"],
                "failedSheets": result["failedSheets"],
            }
        })
    except Exception as e:
        return jsonify({"ok": False, "message": str(e)})



@app.route("/exceptions")
def get_exceptions():
    return jsonify(load_exceptions())


@app.route("/set_exception/<path:product>/<mode>")
def set_exception(product, mode):
    if mode not in ["manufacture", "exclude", "expiry", "clear"]:
        return jsonify({"ok": False, "message": "잘못된 mode입니다."})

    product = clean_text(product)
    if not product:
        return jsonify({"ok": False, "message": "상품명이 없습니다."})

    data = load_exceptions()
    if mode == "clear":
        data.pop(product, None)
    else:
        data[product] = {
            "mode": mode,
            "updatedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        }
    save_exceptions(data)

    try:
        sync_data()
    except Exception:
        pass

    return jsonify({"ok": True, "product": product, "mode": mode})



@app.route("/set_expiry/<path:product>/<expiry_date>")
def set_expiry(product, expiry_date):
    product = clean_text(product)
    expiry_date = clean_text(expiry_date)

    if not product:
        return jsonify({"ok": False, "message": "상품명이 없습니다."})

    try:
        datetime.strptime(expiry_date, "%Y-%m-%d")
    except Exception:
        return jsonify({"ok": False, "message": "유통기한 형식이 올바르지 않습니다. YYYY-MM-DD 형식이어야 합니다."})

    data = load_exceptions()
    data[product] = {
        "mode": "expiry",
        "expiryDate": expiry_date,
        "updatedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    }
    save_exceptions(data)

    try:
        sync_data()
    except Exception:
        pass

    return jsonify({"ok": True, "product": product, "expiryDate": expiry_date})



@app.route("/set_item_expiry/<path:item_key>/<expiry_date>")
def set_item_expiry(item_key, expiry_date):
    item_key = clean_text(item_key)
    expiry_date = clean_text(expiry_date)

    if not item_key:
        return jsonify({"ok": False, "message": "item_key가 없습니다."})

    try:
        datetime.strptime(expiry_date, "%Y-%m-%d")
    except Exception:
        return jsonify({"ok": False, "message": "유통기한 형식이 올바르지 않습니다. YYYY-MM-DD 형식이어야 합니다."})

    data = load_exceptions()
    existing = data.get(item_key, {})
    data[item_key] = {
        **existing,
        "mode": "expiry",
        "expiryDate": expiry_date,
        "scope": "item",
        "updatedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    }
    save_exceptions(data)

    try:
        sync_data()
    except Exception:
        pass

    return jsonify({"ok": True, "itemKey": item_key, "expiryDate": expiry_date})


@app.route("/set_item_exception/<path:item_key>/<mode>")
def set_item_exception(item_key, mode):
    if mode not in ["manufacture", "exclude", "expiry", "clear"]:
        return jsonify({"ok": False, "message": "잘못된 mode입니다."})

    item_key = clean_text(item_key)
    if not item_key:
        return jsonify({"ok": False, "message": "item_key가 없습니다."})

    data = load_exceptions()
    if mode == "clear":
        data.pop(item_key, None)
    else:
        existing = data.get(item_key, {})
        data[item_key] = {
            **existing,
            "mode": mode,
            "scope": "item",
            "updatedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        }
    save_exceptions(data)

    try:
        sync_data()
    except Exception:
        pass

    return jsonify({"ok": True, "itemKey": item_key, "mode": mode})


@app.route("/status")
def status():
    need_sync = True
    updated_at = ""
    count = 0

    if Path(DATA_FILE).exists():
        try:
            with open(DATA_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
            updated_at = data.get("updatedAt", "")
            count = data.get("count", 0)
            dt = datetime.strptime(updated_at, "%Y-%m-%d %H:%M:%S")
            age_minutes = (datetime.now() - dt).total_seconds() / 60
            need_sync = age_minutes >= AUTO_SYNC_MINUTES
        except Exception:
            need_sync = True

    return jsonify({
        "ok": True,
        "needSync": need_sync,
        "updatedAt": updated_at,
        "count": count,
        "autoSyncMinutes": AUTO_SYNC_MINUTES
    })



@app.route("/set_store_item_expiry/<path:item_key>/<store_name>/<expiry_date>")
def set_store_item_expiry(item_key, store_name, expiry_date):
    item_key = clean_text(item_key)
    store_name = clean_text(store_name)
    expiry_date = clean_text(expiry_date)

    if not item_key:
        return jsonify({"ok": False, "message": "item_key가 없습니다."})
    if not store_name:
        return jsonify({"ok": False, "message": "매장명이 없습니다."})

    try:
        datetime.strptime(expiry_date, "%Y-%m-%d")
    except Exception:
        return jsonify({"ok": False, "message": "유통기한 형식이 올바르지 않습니다. YYYY-MM-DD 형식이어야 합니다."})

    # 같은 원본 행이라도 매장별로 다른 유통기한을 저장할 수 있게 매장명을 앞에 붙임
    store_item_key = f"STORE|{store_name}|{item_key}"

    data = load_exceptions()
    existing = data.get(store_item_key, {})
    data[store_item_key] = {
        **existing,
        "mode": "expiry",
        "expiryDate": expiry_date,
        "scope": "store_item",
        "store": store_name,
        "baseItemKey": item_key,
        "updatedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    }
    save_exceptions(data)

    try:
        sync_data()
    except Exception:
        pass

    return jsonify({"ok": True, "itemKey": store_item_key, "store": store_name, "expiryDate": expiry_date})


@app.route("/store_overrides/<store_name>")
def get_store_overrides(store_name):
    store_name = clean_text(store_name)
    data = load_exceptions()
    prefix = f"STORE|{store_name}|"
    result = {}
    for key, value in data.items():
        if key.startswith(prefix):
            base_key = key[len(prefix):]
            result[base_key] = value
    return jsonify({"ok": True, "store": store_name, "overrides": result})



def apply_store_override_to_export(item, store_name, exceptions):
    """엑셀 다운로드 시 선택 매장의 수정 유통기한을 반영한다."""
    row = dict(item)
    store_name = clean_text(store_name)
    if not store_name:
        return row

    override_key = f"STORE|{store_name}|{row.get('itemKey', '')}"
    override = exceptions.get(override_key, {})
    expiry_date = clean_text(override.get("expiryDate", ""))

    if expiry_date:
        row["expiryDate"] = expiry_date
        days, section = section_by_expiry(expiry_date)
        row["daysLeft"] = days
        row["section"] = section
        row["manage"] = "유통기한관리"
        row["storeOverride"] = True

    mode = clean_text(override.get("mode", ""))
    if mode == "manufacture" and not expiry_date:
        row["manage"] = "기한확인필요"
        row["section"] = "기한확인필요"
        row["daysLeft"] = ""
        row["storeOverride"] = True
    elif mode == "exclude" and not expiry_date:
        row["manage"] = "비관리대상"
        row["section"] = "기한관리제외"
        row["daysLeft"] = ""
        row["storeOverride"] = True

    return row


def export_menu_match(item, menu):
    if not menu or menu == "전체상품":
        return True
    if menu == "유통기한관리":
        return item.get("manage") == "유통기한관리"
    if menu == "기한확인필요":
        return item.get("manage") == "기한확인필요"
    if menu == "비관리대상":
        return item.get("manage") == "비관리대상"
    return item.get("section") == menu


def split_export_items(items):
    """형태가 기증/기부이면 기증, 상생/매입이면 상생으로 분리.

    원본에 형태가 비어 있는 상품은 잘못된 시트에 넣거나 누락하지 않고,
    필요할 때만 '형태미확인' 시트로 따로 표시한다.
    """
    grouped = {"기증": [], "상생": []}
    unchecked = []
    for item in items:
        item_type = clean_text(item.get("type", ""))
        if "상생" in item_type or "매입" in item_type:
            grouped["상생"].append(item)
        elif "기증" in item_type or "기부" in item_type:
            grouped["기증"].append(item)
        else:
            unchecked.append(item)
    if unchecked:
        grouped["형태미확인"] = unchecked
    return grouped


def _excel_text(value):
    """Excel에서 허용하지 않는 문자를 제거하고 문자열 길이를 제한."""
    value = clean_text(value)
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]", "", value)[:32767]


def build_grouped_workbook(grouped, headers, make_row, title, summary, widths,
                           highlight_days_col=None, blank_col=None):
    """표준 XLSX 라이브러리로 그룹별 시트가 있는 파일 생성.

    XlsxWriter 우선, 미설치라면 openpyxl 사용. 직접 OOXML 작성은 하지 않는다.
    """
    output = io.BytesIO()
    try:
        import xlsxwriter
    except ImportError:
        xlsxwriter = None

    if xlsxwriter is not None:
        wb = xlsxwriter.Workbook(output, {
            "in_memory": True, "strings_to_formulas": False,
            "strings_to_urls": False,
        })
        title_fmt = wb.add_format({"bold": True, "font_size": 16, "font_color": "#FFFFFF",
                                   "bg_color": "#147D73", "align": "center", "valign": "vcenter"})
        info_fmt = wb.add_format({"bg_color": "#F1F5F9", "font_color": "#334155",
                                  "valign": "vcenter"})
        header_fmt = wb.add_format({"bold": True, "font_color": "#FFFFFF",
                                    "bg_color": "#111827", "border": 1, "align": "center",
                                    "valign": "vcenter", "text_wrap": True})
        text_fmt = wb.add_format({"border": 1, "valign": "top", "text_wrap": True})
        center_fmt = wb.add_format({"border": 1, "align": "center", "valign": "top",
                                    "text_wrap": True})
        warning_fmt = wb.add_format({"border": 1, "align": "center", "valign": "top",
                                     "bg_color": "#FFF2CC", "font_color": "#C65911"})
        for name, sheet_items in grouped.items():
            ws = wb.add_worksheet(name)
            last_col = len(headers) - 1
            ws.merge_range(0, 0, 0, last_col, f"{title} - {name}", title_fmt)
            ws.merge_range(1, 0, 1, last_col, summary(name, len(sheet_items)), info_fmt)
            ws.write_row(2, 0, headers, header_fmt)
            for idx, item in enumerate(sheet_items, start=3):
                values = make_row(item)
                for col, value in enumerate(values):
                    if col == blank_col:
                        ws.write_blank(idx, col, None, center_fmt)
                        continue
                    fmt = warning_fmt if col == highlight_days_col and isinstance(value, (int, float)) and not isinstance(value, bool) and value < 0 else (center_fmt if col in (0, 1, 5, 6, 7, 8) else text_fmt)
                    if isinstance(value, (int, float)) and not isinstance(value, bool):
                        ws.write_number(idx, col, value, fmt)
                    else:
                        ws.write_string(idx, col, _excel_text(value), fmt)
                long_col_texts = [str(v or "") for v in values if isinstance(v, str)]
                longest = max((len(x) for x in long_col_texts), default=0)
                lines = max((len(x.splitlines()) for x in long_col_texts), default=1)
                ws.set_row(idx, min(90, max(28, 18 * max(lines, longest // 35 + 1))))
            for col, width in enumerate(widths):
                ws.set_column(col, col, width)
            ws.set_row(0, 32)
            ws.set_row(1, 25)
            ws.set_row(2, 30)
            ws.freeze_panes(3, 0)
            ws.autofilter(2, 0, max(2, len(sheet_items) + 2), last_col)
            ws.set_landscape()
            ws.set_paper(9)
            ws.fit_to_pages(1, 0)
            ws.hide_gridlines(2)
        wb.close()
        output.seek(0)
        return output

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("엑셀 생성 라이브러리가 없습니다. requirements.txt에 XlsxWriter>=3.1.0을 추가해주세요.") from exc

    wb = Workbook()
    wb.remove(wb.active)
    edge = Side(style="thin", color="CBD5E1")
    for name, sheet_items in grouped.items():
        ws = wb.create_sheet(name)
        last_col = get_column_letter(len(headers))
        ws.merge_cells(f"A1:{last_col}1")
        ws["A1"] = f"{title} - {name}"
        ws.merge_cells(f"A2:{last_col}2")
        ws["A2"] = summary(name, len(sheet_items))
        ws["A1"].font = Font(name="맑은 고딕", size=16, bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", fgColor="147D73")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws["A2"].font = Font(name="맑은 고딕", color="334155")
        ws["A2"].fill = PatternFill("solid", fgColor="F1F5F9")
        ws.append(headers)
        for c in ws[3]:
            c.font = Font(name="맑은 고딕", bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="111827")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = Border(left=edge, right=edge, top=edge, bottom=edge)
        for item in sheet_items:
            values = make_row(item)
            if blank_col is not None:
                values[blank_col] = None
            ws.append([v if isinstance(v, (int, float)) else _excel_text(v) for v in values])
            row = list(ws[ws.max_row])
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
                c.border = Border(left=edge, right=edge, top=edge, bottom=edge)
                if c.data_type == "f":
                    c.data_type = "s"   # 상품명 등 '='로 시작하는 텍스트를 수식으로 처리하지 않음
            if highlight_days_col is not None:
                days = row[highlight_days_col]
                if isinstance(days.value, (int, float)) and days.value < 0:
                    days.fill = PatternFill("solid", fgColor="FFF2CC")
                    days.font = Font(name="맑은 고딕", color="C65911")
            ws.row_dimensions[ws.max_row].height = 30
        for j, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = width
        ws.row_dimensions[1].height = 32
        ws.row_dimensions[2].height = 25
        ws.row_dimensions[3].height = 30
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:{last_col}{max(3, len(sheet_items)+3)}"
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    wb.active = 0
    wb.save(output)
    output.seek(0)
    return output


def build_discount_workbook(grouped, store_label, menu_label):
    """가격 인하 요청서: 기증/상생 시트를 분리하여 다운로드."""
    headers = ["형태", "대분류", "상품명", "가격 원문", "인하 가격", "유통기한", "남은 일수"]
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M")

    def row(item):
        price = clean_text(item.get("priceRaw")) or clean_text(item.get("priceDisplay"))
        return [item.get("type"), item.get("category"), item.get("product"), price,
                "", item.get("expiryDate"), item.get("daysLeft", "")]

    def info(name, count):
        return (f"매장: {store_label} | 분류: {menu_label} | "
                f"{name}: {count:,}건 | 생성: {timestamp}")

    return build_grouped_workbook(grouped, headers, row, "가격인하 취합", info,
                                  [11, 12, 34, 44, 17, 14, 12],
                                  highlight_days_col=6, blank_col=4)


@app.route("/download_excel", methods=["GET", "POST"])
def download_excel():
    """화면 검색 조건을 반영한 가격인하 취합 엑셀(기증/상생 시트)을 다운로드한다."""
    try:
        if not Path(DATA_FILE).exists():
            sync_data()

        with open(DATA_FILE, 'r', encoding='utf-8') as f:
            source = json.load(f)

        items = source.get('items', [])
        exceptions = load_exceptions()
        # 검색 화면(index.html)은 JSON을 POST로 전송하고,
        # 주소 직접 입력/기존 링크는 GET 쿼리 파라미터를 사용한다.
        # 두 방식을 모두 지원하여 다운로드 버튼과 직접 링크가 똑같이 동작하도록 한다.
        if request.method == "POST":
            params = request.get_json(silent=True) or {}
            if not isinstance(params, dict):
                return jsonify({"ok": False, "message": "다운로드 요청 형식이 올바르지 않습니다."}), 400
        else:
            params = request.args

        search_mode = clean_text(params.get('mode', 'product'))
        query = clean_text(params.get('q', ''))
        date_start = clean_text(params.get('start', ''))
        date_end = clean_text(params.get('end', ''))
        menu = clean_text(params.get('menu', '전체상품'))
        store_name = clean_text(params.get('store', ''))

        # 사용자가 화면에서 '가격인하 완료'로 표시한 상품은
        # 상품명+유통기한 조합을 기준으로 요청서에서 제외한다.
        # 기존 브라우저 로컬 저장 키 형식과 동일한 정규화를 사용한다.
        raw_exclude_keys = params.get('excludeKeys', [])
        exclude_keys = set(
            str(key) for key in raw_exclude_keys if isinstance(key, str)
        ) if isinstance(raw_exclude_keys, list) else set()

        def discount_done_key(item):
            name = clean_text(item.get('product', '')).lower().strip()
            name = re.sub(r'[\s\-_.\/·ㆍ()\[\]{}]+', '', name)
            expiry = clean_text(item.get('expiryDate', '')) or '__NO_EXPIRY__'
            return f'{name}|||{expiry}'

        selected = []
        tokens = [x for x in query.lower().split() if x]
        for original in items:
            item = apply_store_override_to_export(original, store_name, exceptions)
            inbound = clean_text(item.get('inboundDate', ''))
            if date_start and inbound < date_start:
                continue
            if date_end and inbound > date_end:
                continue
            if not export_menu_match(item, menu):
                continue
            if query:
                if search_mode == 'box':
                    if clean_text(item.get('boxNo', '')) != query:
                        continue
                else:
                    search_text = clean_text(item.get('search', '')).lower()
                    if not all(token in search_text for token in tokens):
                        continue

            if exclude_keys and discount_done_key(item) in exclude_keys:
                continue

            selected.append(item)

        grouped = split_export_items(selected)
        store_label = store_name or '전체'
        menu_label = menu if menu and menu != '전체상품' else '전체'
        output = build_discount_workbook(grouped, store_label, menu_label)

        filename = f"가격인하_취합_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        response = send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response.headers['Cache-Control'] = 'no-store'
        response.headers['X-Price-Export-Fix'] = 'post-enabled-v7'
        response.headers['X-Price-Export-Sheets'] = '|'.join(grouped.keys())
        return response
    except Exception as e:
        app.logger.exception('가격인하 엑셀 다운로드 실패')
        return jsonify({'ok': False, 'message': f'엑셀 다운로드 처리 오류: {e}'}), 500


# 월별 유통기한 조회는 기존 가격검색·가격인하 다운로드와 독립적으로 작동한다.
EXPIRY_CALENDAR_HTML = '<!doctype html>\n<html lang="ko"><head><meta charset="utf-8"/><meta name="viewport" content="width=device-width,initial-scale=1"/>\n<title>월별 유통기한 조회 - 기빙플러스</title>\n<style>\n:root{font-family:system-ui,-apple-system,"Malgun Gothic",sans-serif;color:#1e293b;background:#f8fafc}\n*{box-sizing:border-box}body{margin:0}main{max-width:1180px;margin:0 auto;padding:30px 18px 64px}\na{color:#147d73;text-decoration:none}a:hover{text-decoration:underline}h1{margin:10px 0;font-size:27px}\n.sub{color:#64748b;margin:6px 0 24px}.panel{background:#fff;border:1px solid #e2e8f0;border-radius:14px;padding:20px;box-shadow:0 2px 8px rgba(0,0,0,.03)}\n.controls{display:flex;align-items:flex-end;gap:12px;flex-wrap:wrap}.control{display:flex;flex-direction:column;gap:6px}\nlabel{font-size:13px;font-weight:700}input,select,button{font:inherit;min-height:42px;border:1px solid #cbd5e1;border-radius:8px;padding:8px 12px;background:white}\nbutton{background:#147d73;color:white;font-weight:700;border-color:#147d73;cursor:pointer}button:disabled{opacity:.55;cursor:default}\n.help{color:#64748b;font-size:13px;margin:14px 0 0}.summary{display:flex;flex-wrap:wrap;gap:12px;margin:18px 0}.stat{background:white;border:1px solid #e2e8f0;border-radius:12px;padding:15px 20px;min-width:160px}\n.stat strong{display:block;font-size:25px;margin-top:4px}.stat small{color:#64748b}\n.table-wrap{overflow:auto;background:white;border:1px solid #e2e8f0;border-radius:14px}table{width:100%;border-collapse:collapse;min-width:900px;font-size:14px}\nth{text-align:left;background:#f1f5f9;white-space:nowrap;position:sticky;top:0}th,td{padding:12px 13px;border-bottom:1px solid #e2e8f0;vertical-align:top}td.product{min-width:210px;max-width:360px;white-space:pre-wrap}\n.empty{padding:48px 20px;text-align:center;color:#64748b}#status{margin:15px 0;color:#475569} .tag{display:inline-block;border-radius:20px;padding:3px 9px;background:#eef2f6;font-size:12px}\n@media(max-width:650px){h1{font-size:23px}main{padding:18px 12px}.panel{padding:15px}}\n</style></head>\n<body><main>\n<a href="/">← 가격검색으로 돌아가기</a>\n<h1>📅 월별 유통기한 조회</h1>\n<p class="sub">가격리스트에 유통기한이 등록된 상품을 원하는 날짜 기준으로 모아 볼 수 있습니다.</p>\n<div class="panel">\n<form id="form" class="controls">\n <div class="control"><label for="target">기준 날짜</label><input id="target" type="date" required value="__DEFAULT_DATE__"/></div>\n <div class="control"><label for="mode">조회 범위</label><select id="mode">\n  <option value="month">선택한 달 전체 (1일~말일)</option>\n  <option value="until">오늘부터 선택 날짜까지</option>\n  <option value="day">선택 날짜 당일만</option>\n </select></div>\n <button type="submit" id="searchBtn">상품 조회</button>\n <button type="button" id="exportBtn" disabled>기증·상생 분리 엑셀 다운로드 (.xlsx)</button>\n <button type="button" id="csvBtn" disabled>엑셀 호환 CSV 다운로드</button>\n</form>\n<p class="help" id="modeHelp">예: 12월 31일을 선택하면 12월 1~31일에 유통기한이 도래하는 상품을 표시합니다.</p>\n</div>\n<div id="status" role="status">상품을 조회하는 중입니다.</div>\n<div class="summary"><div class="stat"><small>조회된 상품 행</small><strong id="total">-</strong></div>\n<div class="stat"><small>기증</small><strong id="donated">-</strong></div>\n<div class="stat"><small>상생</small><strong id="partner">-</strong></div></div>\n<div class="table-wrap"><table><thead><tr><th>유통기한</th><th>형태</th><th>대분류</th><th>상품명</th><th>가격</th><th>남은 일수</th><th>입고일</th><th>매장</th><th>박스번호</th></tr></thead><tbody id="results"><tr><td colspan="9" class="empty">조회 중...</td></tr></tbody></table></div>\n<p class="help">가격리스트에 기한이 기록된 상품 기준입니다. 실제 매장 재고·판매 여부는 별도 확인이 필요하며, 같은 상품이 여러 차례 등록된 경우 여러 행으로 표시될 수 있습니다. · 다운로드 개선 버전 v6 · 기증/상생 시트 분리</p>\n</main>\n<script>\nconst form=document.getElementById(\'form\'), target=document.getElementById(\'target\'), mode=document.getElementById(\'mode\');\nconst help=document.getElementById(\'modeHelp\'), results=document.getElementById(\'results\'), status=document.getElementById(\'status\');\nconst exportBtn=document.getElementById(\'exportBtn\'),csvBtn=document.getElementById(\'csvBtn\');\nlet lastSearchParams=null,lastSearchItems=null,lastSearchMetadata=null;\nconst HELP={month:\'예: 12월 31일 선택 → 12월 1~31일 만료 상품 전체\',until:\'예: 12월 31일 선택 → 오늘부터 12월 31일까지 만료 예정 상품\',day:\'예: 12월 31일 선택 → 12월 31일 당일 만료 상품만\'};\nfunction markResultsStale(){lastSearchParams=null;lastSearchItems=null;lastSearchMetadata=null;exportBtn.disabled=true;csvBtn.disabled=true;}\nmode.addEventListener(\'change\',()=>{help.textContent=HELP[mode.value];markResultsStale();});\ntarget.addEventListener(\'change\',markResultsStale);\nfunction td(value,klass){const cell=document.createElement(\'td\');cell.textContent=value===null||value===undefined?\'\':String(value);if(klass)cell.className=klass;return cell;}\nfunction none(message){results.replaceChildren();const tr=document.createElement(\'tr\');const cell=td(message,\'empty\');cell.colSpan=9;tr.append(cell);results.append(tr);}\nasync function search(){\n const btn=document.getElementById(\'searchBtn\');btn.disabled=true;markResultsStale();status.textContent=\'조회 중입니다...\';none(\'상품을 불러오는 중입니다.\');\n try{\n  const params=new URLSearchParams({date:target.value,mode:mode.value});\n  const response=await fetch(\'/api/expiry_calendar?\'+params.toString(),{cache:\'no-store\'});\n  const data=await response.json();if(!response.ok||!data.ok)throw new Error(data.message||\'조회에 실패했습니다.\');\n  const items=data.items||[];results.replaceChildren();\n  for(const item of items){\n   const tr=document.createElement(\'tr\');\n   const remaining=item.daysLeft;\n   const days=typeof remaining===\'number\'?(remaining<0?Math.abs(remaining)+\'일 지남\':remaining+\'일\'):\'-\';\n   const fields=[item.expiryDate,item.type,item.category,item.product,item.priceDisplay||item.priceRaw,days,item.inboundDate,item.store,item.boxNo];\n   fields.forEach((v,i)=>tr.append(td(v,i===3?\'product\':\'\')));results.append(tr);\n  }\n  if(!items.length)none(\'선택한 기간에 해당하는 유통기한 상품이 없습니다.\');\n  document.getElementById(\'total\').textContent=items.length.toLocaleString(\'ko-KR\');\n  document.getElementById(\'donated\').textContent=items.filter(i=>(i.type||\'\').includes(\'기증\')||(i.type||\'\').includes(\'기부\')).length.toLocaleString(\'ko-KR\');\n  document.getElementById(\'partner\').textContent=items.filter(i=>(i.type||\'\').includes(\'상생\')).length.toLocaleString(\'ko-KR\');\n  status.textContent=data.periodLabel+\' · \'+items.length.toLocaleString(\'ko-KR\')+\'행 · 가격리스트 동기화 \'+(data.updatedAt||\'확인 불가\');\n  lastSearchParams=params.toString();lastSearchItems=items;lastSearchMetadata={periodLabel:data.periodLabel,updatedAt:data.updatedAt};exportBtn.disabled=items.length===0;csvBtn.disabled=items.length===0;\n }catch(e){none(\'조회하지 못했습니다.\');status.textContent=\'오류: \'+e.message;}\n finally{btn.disabled=false;}\n}\nasync function downloadResults(){\n if(!lastSearchParams||!lastSearchItems||lastSearchItems.length===0){\n  status.textContent=\'조회된 상품이 없습니다. 날짜를 다시 선택해 주세요.\';return;\n }\n exportBtn.disabled=true;\n const previousStatus=status.textContent;\n status.textContent=\'조회한 \'+lastSearchItems.length.toLocaleString(\'ko-KR\')+\'건을 엑셀로 만드는 중입니다...\';\n try{\n  const response=await fetch(\'/download_expiry_excel_v6?\'+lastSearchParams,{\n    method:\'POST\',cache:\'no-store\',headers:{\'Content-Type\':\'application/json\'},\n    body:JSON.stringify({items:lastSearchItems,periodLabel:lastSearchMetadata.periodLabel,updatedAt:lastSearchMetadata.updatedAt})\n  });\n  if(!response.ok){let message=\'엑셀 다운로드에 실패했습니다.\';try{const error=await response.json();message=error.message||message;}catch(_e){}throw new Error(message);}\n  const exportVersion=response.headers.get(\'X-Export-Version\');\n  if(exportVersion!==\'xlsx-library-v6\'){\n   throw new Error(\'구버전 서버에서 생성한 엑셀입니다. GitHub 배포 완료 여부를 확인하고 Ctrl+F5로 새로고침해주세요.\');\n  }\n  const sheetHeader=response.headers.get(\'X-Export-Sheets\')||\'\';\n  const actualSheets=sheetHeader.split(\'|\');\n  if(!actualSheets.includes(\'기증\')||!actualSheets.includes(\'상생\')){\n   throw new Error(\'다운로드 서버에서 기증/상생 시트를 확인할 수 없습니다. v6 배포 상태를 확인해주세요.\');\n  }\n  const reportedCount=Number(response.headers.get(\'X-Export-Row-Count\'));\n  if(!Number.isFinite(reportedCount)||reportedCount!==lastSearchItems.length){\n   throw new Error(\'조회 결과 수와 엑셀 생성 수가 다릅니다. 화면 새로고침 후 다시 시도해 주세요.\');\n  }\n  const blob=await response.blob();\n  if(blob.size<1000)throw new Error(\'다운로드된 엑셀 파일이 비정상적으로 작습니다.\');\n  const url=URL.createObjectURL(blob);\n  const link=document.createElement(\'a\');link.href=url;\n  link.download=\'유통기한_기증상생_v6_\'+target.value+\'_\'+mode.value+\'_\'+reportedCount+\'건.xlsx\';\n  document.body.appendChild(link);link.click();link.remove();\n  setTimeout(()=>URL.revokeObjectURL(url),60000);\n  status.textContent=previousStatus+\' · \'+reportedCount.toLocaleString(\'ko-KR\')+\'건 · \'+sheetHeader.replaceAll(\'|\',\' / \')+\' 시트로 다운로드 완료\';\n }catch(e){status.textContent=\'엑셀 다운로드 오류: \'+e.message;}\n finally{exportBtn.disabled=!lastSearchItems||lastSearchItems.length===0;}\n}\nfunction downloadCsvResults(){\n if(!lastSearchItems||lastSearchItems.length===0){status.textContent=\'조회된 상품이 없습니다.\';return;}\n const headings=[\'유통기한\',\'형태\',\'대분류\',\'상품명\',\'가격\',\'남은 일수\',\'입고일\',\'매장\',\'박스번호\'];\n const fields=[\'expiryDate\',\'type\',\'category\',\'product\',\'priceDisplay\',\'daysLeft\',\'inboundDate\',\'store\',\'boxNo\'];\n function csvCell(value){\n  let v=String(value===undefined||value===null?\'\':value);\n  if(/^[\\s]*[=+@]/.test(v)||(/^[\\s]*-/.test(v) && !/^\\s*-\\d+(?:\\.\\d+)?\\s*$/.test(v)))v=String.fromCharCode(39)+v;\n  return String.fromCharCode(34)+v.replace(/"/g,\'""\')+String.fromCharCode(34);\n }\n const lines=[headings.map(csvCell).join(\',\')];\n for(const item of lastSearchItems){\n  const record=fields.map(f=>f===\'priceDisplay\'?(item.priceDisplay||item.priceRaw||\'\'):item[f]);\n  lines.push(record.map(csvCell).join(\',\'));\n }\n const blob=new Blob([String.fromCharCode(0xFEFF)+lines.join(\'\\r\\n\')],{type:\'text/csv;charset=utf-8;\'});\n const url=URL.createObjectURL(blob),link=document.createElement(\'a\');\n link.href=url;link.download=\'유통기한조회_\'+target.value+\'_\'+mode.value+\'_\'+lastSearchItems.length+\'건.csv\';\n document.body.appendChild(link);link.click();link.remove();\n setTimeout(()=>URL.revokeObjectURL(url),60000);\n status.textContent=\'조회된 \'+lastSearchItems.length.toLocaleString(\'ko-KR\')+\'건을 CSV로 저장했습니다. Excel에서 열 수 있습니다.\';\n}\ncsvBtn.addEventListener(\'click\',downloadCsvResults);\nexportBtn.addEventListener(\'click\',downloadResults);\nform.addEventListener(\'submit\',event=>{event.preventDefault();search();});\nsearch();\n</script></body></html>'


def get_expiry_calendar_rows(target_date, mode):
    if not Path(DATA_FILE).exists():
        sync_data()
    with open(DATA_FILE, "r", encoding="utf-8") as f:
        source = json.load(f)
    today = date.today()
    start = target_date.replace(day=1)
    if mode == "month":
        end_day = calendar.monthrange(target_date.year, target_date.month)[1]
        end = target_date.replace(day=end_day)
        period_label = f"{start:%Y-%m-%d} ~ {end:%Y-%m-%d}"
    elif mode == "until":
        start, end = today, target_date
        period_label = f"{start:%Y-%m-%d} ~ {end:%Y-%m-%d}"
    else:
        start = end = target_date
        period_label = f"{target_date:%Y-%m-%d} 당일"

    filtered = []
    for item in source.get("items", []):
        product = clean_text(item.get("product", ""))
        if "전량소진" in product:
            continue
        expiry_text = clean_text(item.get("expiryDate", ""))
        try:
            expiry = date.fromisoformat(expiry_text)
        except ValueError:
            continue
        if start <= expiry <= end and item.get("manage") == "유통기한관리":
            filtered.append({
                "type": item.get("type", ""),
                "category": item.get("category", ""),
                "product": product,
                "priceDisplay": item.get("priceDisplay", ""),
                "priceRaw": item.get("priceRaw", ""),
                "inboundDate": item.get("inboundDate", ""),
                "expiryDate": expiry_text,
                "daysLeft": (expiry-today).days,
                "store": item.get("store", ""),
                "boxNo": item.get("boxNo", ""),
            })
    filtered.sort(key=lambda x: (x["expiryDate"], x["product"], x["type"]))
    return filtered, source.get("updatedAt", ""), period_label


@app.route("/expiry_calendar")
def expiry_calendar():
    today = date.today()
    last = today.replace(day=calendar.monthrange(today.year, today.month)[1])
    response = Response(EXPIRY_CALENDAR_HTML.replace("__DEFAULT_DATE__", last.isoformat()),
                        mimetype="text/html")
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response.headers["X-Expiry-App-Version"] = "xlsx-library-v6"
    return response


@app.route("/api/expiry_calendar")
def api_expiry_calendar():
    date_text = clean_text(request.args.get("date", ""))
    mode = clean_text(request.args.get("mode", "month"))
    try:
        target_date = date.fromisoformat(date_text)
    except ValueError:
        return jsonify({"ok": False, "message": "날짜를 YYYY-MM-DD 형식으로 입력해주세요."}), 400
    if mode not in ("month", "until", "day"):
        return jsonify({"ok": False, "message": "올바르지 않은 조회 범위입니다."}), 400
    try:
        items, updated_at, period_label = get_expiry_calendar_rows(target_date, mode)
    except Exception as e:
        return jsonify({"ok": False, "message": str(e)}), 500
    return jsonify({"ok": True, "items": items, "count": len(items),
                    "periodLabel": period_label, "updatedAt": updated_at})




def build_expiry_calendar_workbook(items, period_label, updated_at):
    """월별 유통기한 조회 결과를 형태별 시트로 분리한 Excel 다운로드."""
    grouped = split_export_items(items)
    headers = ["유통기한", "형태", "대분류", "상품명", "가격", "남은 일수", "입고일", "매장", "박스번호"]

    def row(item):
        price = clean_text(item.get("priceDisplay")) or clean_text(item.get("priceRaw"))
        return [item.get("expiryDate"), item.get("type"), item.get("category"),
                item.get("product"), price, item.get("daysLeft", ""),
                item.get("inboundDate"), item.get("store"), item.get("boxNo")]

    def info(name, count):
        return (f"기간: {period_label} | {name} {count:,}건 | "
                f"전체 {len(items):,}건 | 원본 동기화: {updated_at or '확인 불가'}")

    return build_grouped_workbook(grouped, headers, row, "유통기한 조회 결과", info,
                                  [16, 12, 15, 43, 38, 15, 16, 21, 14],
                                  highlight_days_col=5)


@app.route("/download_expiry_excel_v6", methods=["GET", "POST"])
@app.route("/download_expiry_excel", methods=["GET", "POST"])
def download_expiry_excel():
    """조회 화면에 표시된 목록을 그대로 Excel로 내보낸다.

    POST는 화면의 조회 결과를 사용하므로 재조회 시점/조건 차이로
    조회 화면과 엑셀 목록이 달라지는 문제를 방지한다.
    기존 GET 방식도 호환을 위해 계속 지원한다.
    """
    date_text = clean_text(request.args.get("date", ""))
    mode = clean_text(request.args.get("mode", "month"))
    try:
        target_date = date.fromisoformat(date_text)
    except ValueError:
        return jsonify({"ok": False, "message": "날짜를 YYYY-MM-DD 형식으로 입력해주세요."}), 400
    if mode not in ("month", "until", "day"):
        return jsonify({"ok": False, "message": "올바르지 않은 조회 범위입니다."}), 400

    try:
        if request.method == "POST":
            payload = request.get_json(silent=True)
            if not isinstance(payload, dict) or not isinstance(payload.get("items"), list):
                return jsonify({"ok": False, "message": "화면의 상품 목록을 전달받지 못했습니다. 다시 조회해주세요."}), 400
            raw_items = payload["items"]
            if len(raw_items) > 30000:
                return jsonify({"ok": False, "message": "한 번에 다운로드할 수 있는 최대 30,000건을 초과했습니다."}), 413
            allowed_fields = ("expiryDate", "type", "category", "product", "priceDisplay", "priceRaw",
                              "daysLeft", "inboundDate", "store", "boxNo")
            if not all(isinstance(item, dict) for item in raw_items):
                return jsonify({"ok": False, "message": "조회 데이터 형식이 올바르지 않습니다."}), 400
            items = [{key: item.get(key, "") for key in allowed_fields} for item in raw_items]
            period_label = clean_text(payload.get("periodLabel", ""))[:120]
            updated_at = clean_text(payload.get("updatedAt", ""))[:40]
        else:
            items, updated_at, period_label = get_expiry_calendar_rows(target_date, mode)

        if not items:
            return jsonify({"ok": False, "message": "선택한 기간에 유통기한 상품이 0건입니다. 날짜를 확인한 뒤 다시 조회해주세요."}), 422

        output = build_expiry_calendar_workbook(items, period_label, updated_at)
        from zipfile import ZipFile
        from xml.etree import ElementTree
        grouped = split_export_items(items)
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        output.seek(0)
        with ZipFile(output) as archive:
            if archive.testzip() is not None:
                raise ValueError("엑셀 압축 파일 검사 실패")
            workbook_root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            sheet_names = [element.get("name") for element in workbook_root.findall("s:sheets/s:sheet", ns)]
            if sheet_names != list(grouped):
                raise ValueError(f"엑셀 시트 검사 실패: {sheet_names}")
            for sheet_number, (sheet_name, group_items) in enumerate(grouped.items(), start=1):
                root = ElementTree.fromstring(archive.read(f"xl/worksheets/sheet{sheet_number}.xml"))
                rows = root.findall("s:sheetData/s:row", ns)
                if len(rows) != len(group_items) + 3:
                    raise ValueError(f"{sheet_name} 시트 행 수 검사 실패: 예상 {len(group_items)}건")
            if sum(len(group) for group in grouped.values()) != len(items):
                raise ValueError("일부 상품이 엑셀에서 누락되었습니다.")
        output.seek(0)

        label = {"month": "월별", "until": "기준일까지", "day": "당일"}[mode]
        filename = f"유통기한_기증상생_v6_{target_date:%Y%m%d}_{label}_{len(items)}건.xlsx"
        response = send_file(output, as_attachment=True, download_name=filename,
                             mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response.headers["Cache-Control"] = "no-store"
        response.headers["X-Export-Row-Count"] = str(len(items))
        response.headers["X-Export-Sheets"] = "|".join(grouped.keys())
        response.headers["X-Export-Version"] = "xlsx-library-v6"
        return response
    except Exception as e:
        app.logger.exception("유통기한 조회 엑셀 다운로드 실패")
        return jsonify({"ok": False, "message": f"엑셀 다운로드 처리 오류: {e}"}), 500


@app.route("/health")
def health():
    return "ok"


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "8000"))
    app.run(host="0.0.0.0", port=port)
